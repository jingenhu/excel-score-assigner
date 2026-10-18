# -*- coding: utf-8 -*-
"""
score_assigner.py — 教务成绩自动赋值工具

根据「目标总评成绩」反向推导出满足取值区间的平时分，仅写入赋值列，
完整保留原表的公式、单元格格式与表格结构。

默认用法：
    python score_assigner.py

可选参数（覆盖顶部配置）：
    python score_assigner.py --file scores.xlsx --sheet Sheet1
    python score_assigner.py --dry-run          # 只试算，不写文件

作者：胡金根
许可：MIT
"""

from __future__ import annotations

import argparse
import math
import os
import random
import shutil
import sys
from collections import Counter, deque
from datetime import datetime
from typing import Iterable, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# ==========================================================================
# 顶部配置区：切换学科 / 课程，只改这里
# ==========================================================================

# ---- 基础文件配置 ----
FILE_PATH = os.path.join("examples", "demo_scores.xlsx")  # 目标 Excel 文件路径
SHEET_NAME = "Sheet1"                                     # 工作表名称
DATA_START_ROW = 2                                        # 数据起始行号（跳过表头）

# ---- 列映射（字母列号）----
ID_COL = "B"          # 学号列（只读，用于报告定位）
NAME_COL = "C"        # 姓名列（只读，用于报告定位）
ATTEND_COL = "H"      # 出勤成绩列（只读）
HW_COL = "J"          # 平时分列（写入列）
TARGET_COL = "L"      # 目标总评列（只读）

# ---- 赋值列规则 ----
HW_MIN = 60           # 平时分最小值
HW_MAX = 99           # 平时分最大值
WEIGHT_I = 0.8        # 平时成绩权重
WEIGHT_J = 0.2        # 出勤成绩权重
J_VALUE = 100         # 出勤列固定值（当出勤单元格为空时使用）

# ---- 写入策略 ----
MAX_CONSECUTIVE = 2   # 同一数值最多连续出现几次（防连续重复校验）
SEED = 20260101       # 随机种子，保证同配置多次运行结果可复现
OUTPUT_SUFFIX = "_filled"   # 输出文件后缀

# ==========================================================================


def _force_utf8_stdout() -> None:
    """兼容 Windows 中文 GBK 控制台：无法输出 UTF-8 时退回 replace，不抛异常。"""
    for stream in (sys.stdout, sys.stderr):
        enc = (getattr(stream, "encoding", "") or "").lower()
        if enc not in ("utf-8", "utf8"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
            except Exception:
                pass


def excel_round(value: float, digits: int = 0) -> float:
    """
    与 Excel ROUND 一致的「半进位（远离零）」取整。

    注意：Python 内置 round() 是银行家舍入（round-half-to-even），
    round(90.5) == 90 而 Excel ROUND(90.5,0) == 91。
    若直接用 round() 反解，会在 .5 边界上系统性地错位。
    """
    factor = 10 ** digits
    scaled = value * factor
    if scaled >= 0:
        return math.floor(scaled + 0.5) / factor
    return math.ceil(scaled - 0.5) / factor


def evaluate(hw: float, attend: float) -> int:
    """按配置的权重链计算总评：ROUND(平时分 × WEIGHT_I + 出勤 × WEIGHT_J)。"""
    return int(excel_round(WEIGHT_I * hw + WEIGHT_J * attend))


def candidates_for(target: int, attend: float) -> List[int]:
    """
    反向推导：返回所有能让总评恰好等于 target 的整数平时分。

    推导过程（设 x 为平时分，a 为出勤分）：
        ROUND(w_i·x + w_j·a) == T
      ⟺ T - 0.5 ≤ w_i·x + w_j·a < T + 0.5
      ⟺ (T - 0.5 - w_j·a) / w_i ≤ x < (T + 0.5 - w_j·a) / w_i
    取该半开区间与 [HW_MIN, HW_MAX] 的交集内的全部整数。
    """
    return [hw for hw in range(HW_MIN, HW_MAX + 1) if evaluate(hw, attend) == target]


class Assigner:
    """
    候选值选取器：在可达集合内做「轮转 + 使用次数最少优先」，
    同时抑制同一数值的连续出现，使结果既落在合法区间内，又保持合理的离散度。
    """

    def __init__(self, seed: int = SEED, max_consecutive: int = MAX_CONSECUTIVE):
        self.rng = random.Random(seed)
        self.usage: Counter = Counter()
        self.recent: deque = deque(maxlen=max(1, max_consecutive))
        self.relaxed = 0  # 因候选集过小而放宽防重复约束的次数

    def pick(self, candidates: Sequence[int]) -> int:
        pool: List[int] = list(candidates)
        if not pool:
            raise ValueError("candidates 不能为空")

        # 防连续重复校验：若再选这个值会连续重复超限，则先行剔除
        if len(pool) > 1 and len(self.recent) == self.recent.maxlen:
            fresh = [c for c in pool if not all(v == c for v in self.recent)]
            if fresh:
                pool = fresh
            else:
                self.relaxed += 1  # 候选集本身太小，只能放宽

        # 离散度校验：优先选使用次数最少的候选，避免结果集中到单一数值
        least = min(self.usage[c] for c in pool)
        pool = [c for c in pool if self.usage[c] == least]

        value = self.rng.choice(pool)
        self.usage[value] += 1
        self.recent.append(value)
        return value


def _cell_text(ws, row: int, col_letter: str) -> str:
    try:
        v = ws.cell(row=row, column=column_index_from_string(col_letter)).value
    except Exception:
        return ""
    return "" if v is None else str(v).strip()


def _save_with_retry(wb, out_path: str, attempts: int = 3) -> str:
    """
    文件占用兼容：目标文件被 Excel / WPS 打开时会抛 PermissionError，
    重试若干次后改用带时间戳的文件名，不中断流程。
    """
    target = out_path
    for i in range(attempts):
        try:
            wb.save(target)
            return target
        except PermissionError:
            if i == attempts - 1:
                stem, ext = os.path.splitext(out_path)
                target = f"{stem}_{datetime.now().strftime('%H%M%S')}{ext}"
                wb.save(target)
                return target
            import time

            time.sleep(0.8)
    return target


def run(
    file_path: str,
    sheet_name: str = SHEET_NAME,
    dry_run: bool = False,
    verbose: bool = True,
) -> dict:
    """执行一次赋值流程，返回结构化结果（便于测试与二次调用）。"""
    src = os.path.abspath(file_path)
    if not os.path.exists(src):
        raise FileNotFoundError(f"找不到文件：{src}")

    # 1) 双开工作簿：data_only 版本读公式缓存值，原始版本负责写入以保留公式
    wb_values = load_workbook(src, data_only=True)
    wb = load_workbook(src)

    if sheet_name not in wb.sheetnames:
        raise KeyError(f"工作簿中没有名为 {sheet_name!r} 的工作表，现有：{wb.sheetnames}")

    ws = wb[sheet_name]
    ws_v = wb_values[sheet_name]

    hw_idx = column_index_from_string(HW_COL)
    att_idx = column_index_from_string(ATTEND_COL)
    tgt_idx = column_index_from_string(TARGET_COL)

    assigner = Assigner()
    assigned: List[dict] = []
    unreachable: List[dict] = []
    skipped = 0

    # 2) 逐行反解并写入
    for row in range(DATA_START_ROW, ws.max_row + 1):
        raw_target = ws_v.cell(row=row, column=tgt_idx).value

        # 空值行自动跳过
        if raw_target is None or str(raw_target).strip() == "":
            skipped += 1
            continue

        try:
            target = int(round(float(raw_target)))
        except (TypeError, ValueError):
            unreachable.append({"row": row, "target": raw_target, "reason": "目标值非数值"})
            continue

        raw_attend = ws_v.cell(row=row, column=att_idx).value
        if raw_attend is None or str(raw_attend).strip() == "":
            attend = float(J_VALUE)  # 边界值兜底：出勤列空 → 取固定值
        else:
            try:
                attend = float(raw_attend)
            except (TypeError, ValueError):
                attend = float(J_VALUE)

        # 出勤本身也参与加权，必须先取整到 Excel 语义再反解
        attend = excel_round(attend)

        cands = candidates_for(target, attend)
        if not cands:
            unreachable.append(
                {"row": row, "target": target, "reason": "目标在合法区间内不可达"}
            )
            continue

        hw = assigner.pick(cands)
        ws.cell(row=row, column=hw_idx).value = hw
        assigned.append(
            {
                "row": row,
                "id": _cell_text(ws, row, ID_COL),
                "name": _cell_text(ws, row, NAME_COL),
                "attend": attend,
                "hw": hw,
                "target": target,
                "verified": evaluate(hw, attend) == target,
            }
        )

    # 3) 打开时强制重算，避免公式缓存值滞后
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass

    ok = sum(1 for a in assigned if a["verified"])

    # 4) 落盘（自动备份原表）
    out_path: Optional[str] = None
    backup_path: Optional[str] = None
    if not dry_run and assigned:
        stem, ext = os.path.splitext(src)
        backup_path = f"{stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_backup{ext}"
        shutil.copy2(src, backup_path)

        out_path = _save_with_retry(wb, f"{stem}{OUTPUT_SUFFIX}{ext}")

    wb.close()
    wb_values.close()

    values = [a["hw"] for a in assigned]
    mean = sum(values) / len(values) if values else 0.0
    std = (sum((v - mean) ** 2 for v in values) / len(values)) ** 0.5 if values else 0.0

    result = {
        "file": src,
        "output": out_path,
        "backup": backup_path,
        "rows_scanned": assigned.__len__() + len(unreachable),
        "assigned": len(assigned),
        "verified": ok,
        "match_rate": (ok / len(assigned) * 100) if assigned else 0.0,
        "skipped_empty": skipped,
        "unreachable": unreachable,
        "hw_min": min(values) if values else None,
        "hw_max": max(values) if values else None,
        "hw_mean": round(mean, 2),
        "hw_std": round(std, 2),
        "relaxed": assigner.relaxed,
        "dry_run": dry_run,
    }

    if verbose:
        _print_report(result)
    return result


def _print_report(r: dict) -> None:
    line = "-" * 62
    print(line)
    print("成绩赋值结果")
    print(line)
    print(f"源文件      : {r['file']}")
    print(f"输出文件    : {r['output'] or '（试算模式，未写入）'}")
    print(f"原表备份    : {r['backup'] or '（试算模式，未备份）'}")
    print(line)
    print(f"扫描行数    : {r['rows_scanned'] + r['skipped_empty']}")
    print(f"赋值成功    : {r['assigned']}")
    print(f"校验通过    : {r['verified']}  ({r['match_rate']:.2f}%)")
    print(f"空行跳过    : {r['skipped_empty']}")
    print(f"目标不可达  : {len(r['unreachable'])}")
    if r["hw_min"] is not None:
        print(f"平时分区间  : {r['hw_min']} ~ {r['hw_max']}   均值 {r['hw_mean']}   标准差 {r['hw_std']}")
    print(f"放宽防重复  : {r['relaxed']} 次（候选集过小时触发）")

    if r["unreachable"]:
        print(line)
        print("以下行未赋值，请人工确认：")
        for u in r["unreachable"][:20]:
            print(f"  第 {u['row']} 行  目标={u['target']}  {u['reason']}")
        if len(r["unreachable"]) > 20:
            print(f"  … 另有 {len(r['unreachable']) - 20} 行")
    print(line)


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="教务成绩自动赋值工具")
    p.add_argument("--file", default=None, help=f"目标 Excel 路径，默认 {FILE_PATH}")
    p.add_argument("--sheet", default=None, help=f"工作表名，默认 {SHEET_NAME}")
    p.add_argument("--dry-run", action="store_true", help="只试算不写文件")
    return p.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    _force_utf8_stdout()
    args = parse_args(argv)
    try:
        run(
            file_path=args.file or FILE_PATH,
            sheet_name=args.sheet or SHEET_NAME,
            dry_run=args.dry_run,
        )
    except Exception as exc:  # 顶层兜底：不把堆栈抛给非技术用户
        print(f"[错误] {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
