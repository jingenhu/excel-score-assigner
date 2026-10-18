# -*- coding: utf-8 -*-
"""单元测试：python -m unittest discover -s tests -v"""

import os
import shutil
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import score_assigner as sa  # noqa: E402


def _make_sheet(path, rows):
    """按 score_assigner 的列布局造一张最小成绩表：H=出勤, J=平时分, K=总评(公式), L=目标。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i, (attend, target) in enumerate(rows, start=2):
        ws[f"B{i}"] = f"2023{i:04d}"
        ws[f"C{i}"] = f"学生{i}"
        ws[f"H{i}"] = attend
        ws[f"J{i}"] = None
        ws[f"K{i}"] = f"=ROUND(J{i}*0.8+H{i}*0.2,0)"
        ws[f"L{i}"] = target
    wb.save(path)


class TestRounding(unittest.TestCase):
    def test_half_up_matches_excel(self):
        # Excel ROUND 半进位；Python round 是银行家舍入
        self.assertEqual(sa.excel_round(90.5), 91)
        self.assertEqual(sa.excel_round(89.5), 90)
        self.assertEqual(round(90.5), 90)  # 证明两者确有差异
        self.assertEqual(sa.excel_round(90.4), 90)
        self.assertEqual(sa.excel_round(90.6), 91)


class TestCandidates(unittest.TestCase):
    def test_known_case(self):
        # 0.8x + 20 ∈ [89.5, 90.5) → x ∈ [86.875, 88.125) → {87, 88}
        self.assertEqual(sa.candidates_for(90, 100), [87, 88])

    def test_all_candidates_hit_target(self):
        for target in range(70, 99):
            for attend in (85, 90, 95, 100):
                for hw in sa.candidates_for(target, attend):
                    self.assertEqual(sa.evaluate(hw, attend), target)
                    self.assertGreaterEqual(hw, sa.HW_MIN)
                    self.assertLessEqual(hw, sa.HW_MAX)

    def test_unreachable(self):
        # 出勤 100 时总评下限为 68，目标 50 不可能达成
        self.assertEqual(sa.candidates_for(50, 100), [])


class TestEndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_full_run_is_exact_and_keeps_formulas(self):
        path = os.path.join(self.tmp, "scores.xlsx")
        _make_sheet(path, [(100, 90), (95, 85), (90, 78), (85, 70)] * 5)

        res = sa.run(path, verbose=False)

        self.assertEqual(res["assigned"], 20)
        self.assertEqual(res["verified"], 20)
        self.assertEqual(res["match_rate"], 100.0)
        self.assertTrue(os.path.exists(res["backup"]), "必须自动备份原表")

        wb = load_workbook(res["output"])
        ws = wb["Sheet1"]
        for r in range(2, 22):
            self.assertEqual(ws[f"K{r}"].value, f"=ROUND(J{r}*0.8+H{r}*0.2,0)")
            self.assertIsInstance(ws[f"J{r}"].value, int)
        wb.close()

    def test_unreachable_target_is_reported_not_crashing(self):
        path = os.path.join(self.tmp, "bad.xlsx")
        _make_sheet(path, [(100, 90), (100, 50)])

        res = sa.run(path, verbose=False)

        self.assertEqual(res["assigned"], 1)
        self.assertEqual(len(res["unreachable"]), 1)
        self.assertEqual(res["unreachable"][0]["row"], 3)

    def test_empty_target_rows_are_skipped(self):
        path = os.path.join(self.tmp, "sparse.xlsx")
        _make_sheet(path, [(100, 90), (95, None), (90, 80)])

        res = sa.run(path, verbose=False)

        self.assertEqual(res["assigned"], 2)
        self.assertEqual(res["skipped_empty"], 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
