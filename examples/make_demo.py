# -*- coding: utf-8 -*-
"""
make_demo.py — 生成示例成绩表 examples/demo_scores.xlsx

表结构（12 列，与 score_assigner.py 顶部配置一一对应）：
    A 序号  B 学号  C 姓名  D 年级  E 专业班级  F 课程名称
    G 学分  H 出勤成绩  I 作业成绩  J 平时分（待赋值）  K 总评（公式）  L 目标总评

K 列为公式列，用于验证工具「只写赋值列、保留原有公式」的行为。
姓名与学号均为虚构数据。
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo_scores.xlsx")

HEADERS = [
    ("A", "序号", 6),
    ("B", "学号", 14),
    ("C", "姓名", 10),
    ("D", "年级", 8),
    ("E", "专业班级", 16),
    ("F", "课程名称", 20),
    ("G", "学分", 7),
    ("H", "出勤成绩", 10),
    ("I", "作业成绩", 10),
    ("J", "平时分", 10),
    ("K", "总评", 10),
    ("L", "目标总评", 10),
]

SURNAME = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜"
GIVEN = ["伟", "芳", "娜", "敏", "静", "强", "磊", "洋", "艳", "勇",
         "军", "杰", "娟", "涛", "明", "超", "秀英", "霞", "平", "刚"]
ATTEND = [100, 100, 95, 90, 100, 85, 95, 100]


def build() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    center = Alignment(horizontal="center", vertical="center")

    for col, title, width in HEADERS:
        cell = ws[f"{col}1"]
        cell.value = title
        cell.font = Font(name="微软雅黑", size=10, bold=True)
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[col].width = width

    for i in range(60):
        r = i + 2
        name = SURNAME[i % len(SURNAME)] + GIVEN[i % len(GIVEN)]
        attend = ATTEND[i % len(ATTEND)]
        # 目标总评落在可达区间内（出勤 85~100 时可达下限约 65，上限约 99）
        target = 70 + (i * 7) % 27

        ws[f"A{r}"] = i + 1
        ws[f"B{r}"] = f"2023{i // 40 + 1:02d}{i + 1:04d}"
        ws[f"C{r}"] = name
        ws[f"D{r}"] = "2023级"
        ws[f"E{r}"] = "统计学1班" if i % 2 == 0 else "统计学2班"
        ws[f"F{r}"] = "应用统计学"
        ws[f"G{r}"] = 3
        ws[f"H{r}"] = attend
        ws[f"I{r}"] = 70 + (i * 3) % 25
        ws[f"J{r}"] = None                      # 平时分：留给工具写入
        ws[f"K{r}"] = f"=ROUND(J{r}*0.8+H{r}*0.2,0)"   # 总评：公式列，工具不得覆盖
        ws[f"L{r}"] = target

        for col, _, _ in HEADERS:
            c = ws[f"{col}{r}"]
            c.font = Font(name="微软雅黑", size=10)
            c.alignment = center
            c.border = border

    ws.freeze_panes = "A2"
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"已生成示例表：{path}  （60 行数据）")
